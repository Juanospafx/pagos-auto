import re
import unicodedata
from difflib import SequenceMatcher

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


# Archivos del periodo. Solo es necesario cambiar estas rutas cada mes.
ARCHIVO_NOMINA = "Nomina Julio.xlsm"
HOJA_NOMINA = "NOMINA"
ARCHIVO_BASE = "reporte_revision_pbs Julio 2026.xlsm"
ARCHIVO_SALIDA = "retiradas.xlsx"
TITULO_PERIODO = "RETIRADAS - JULIO 2026"


def texto_comparable(valor):
    if pd.isna(valor):
        return ""
    texto = unicodedata.normalize("NFKD", str(valor).upper())
    texto = "".join(
        caracter
        for caracter in texto
        if not unicodedata.combining(caracter)
    )
    return re.sub(r"[^A-Z0-9]", "", texto)


def normalizar_cedula(valor):
    texto = texto_comparable(valor)
    if re.fullmatch(r"\d+\.0", "" if pd.isna(valor) else str(valor).strip()):
        texto = texto[:-1]
    return texto.zfill(11) if texto.isdigit() else texto


def normalizar_documento(valor):
    texto = texto_comparable(valor)
    if texto in {"", "0", "00", "00000000000", "NAN", "NONE"}:
        return ""
    original = "" if pd.isna(valor) else str(valor).strip()
    if re.fullmatch(r"\d+\.0", original):
        texto = texto[:-1]
    return texto.zfill(11) if texto.isdigit() else texto


def normalizar_solicitud(valor):
    solicitud = texto_comparable(valor)
    # La misma solicitud puede aparecer como SCOT o CREC.
    return re.sub(r"^SCOT", "CREC", solicitud)


def preparar_nombre(nombre):
    if pd.isna(nombre):
        return "", ""
    texto = unicodedata.normalize("NFKD", str(nombre).upper())
    texto = "".join(
        caracter
        for caracter in texto
        if not unicodedata.combining(caracter)
    )
    palabras = re.findall(r"[A-Z0-9]+", texto)
    return "".join(palabras), "".join(sorted(palabras))


def similitud_nombres(nombre_1, nombre_2):
    directo_1, ordenado_1 = preparar_nombre(nombre_1)
    directo_2, ordenado_2 = preparar_nombre(nombre_2)
    if not directo_1 or not directo_2:
        return 0.0
    return max(
        SequenceMatcher(None, directo_1, directo_2).ratio(),
        SequenceMatcher(None, ordenado_1, ordenado_2).ratio(),
    )


def edad_cumplida(fecha_nacimiento, fecha_evaluacion):
    if pd.isna(fecha_nacimiento) or pd.isna(fecha_evaluacion):
        return np.nan
    edad = fecha_evaluacion.year - fecha_nacimiento.year
    if (fecha_evaluacion.month, fecha_evaluacion.day) < (
        fecha_nacimiento.month,
        fecha_nacimiento.day,
    ):
        edad -= 1
    return edad


def sumar_anos(fecha, anos):
    if pd.isna(fecha):
        return pd.NaT
    return fecha + pd.DateOffset(years=anos)


def nombre_beneficiario_base(df):
    return (
        df[
            [
                "BEN_PRIMER_NOMBRE",
                "BEN_SEGUNDO_NOMBRE",
                "BEN_PRIMER_APELLIDO",
                "BEN_SEGUNDO_APELLIDO",
            ]
        ]
        .fillna("")
        .astype(str)
        .agg(" ".join, axis=1)
        .str.upper()
        .str.strip()
        .str.replace(r"\s+", " ", regex=True)
    )


def seleccionar_registro_base(fila_nomina, base):
    candidatos = base.loc[
        base["_CEDULA"].eq(fila_nomina["_CEDULA"])
        & base["_SOLICITUD"].eq(fila_nomina["_SOLICITUD"])
    ].copy()
    if candidatos.empty:
        return None, "SIN COINCIDENCIA"

    documento = fila_nomina["_DOCUMENTO_BEN"]
    if documento:
        por_documento = candidatos.loc[
            candidatos["_DOCUMENTO_BEN"].eq(documento)
        ]
        if not por_documento.empty:
            return por_documento.iloc[0], "DOCUMENTO"

    nombre = fila_nomina["_NOMBRE_BEN"]
    por_nombre = candidatos.loc[candidatos["_NOMBRE_BEN"].eq(nombre)]
    if not por_nombre.empty:
        return por_nombre.iloc[0], "NOMBRE"

    candidatos["_SIMILITUD"] = candidatos["NOMBRE_BENEFICIARIO"].map(
        lambda valor: similitud_nombres(
            fila_nomina["NOMBRES_BENEFICIARIO"], valor
        )
    )
    candidatos.sort_values("_SIMILITUD", ascending=False, inplace=True)
    mejor = candidatos.iloc[0]
    segundo_puntaje = (
        candidatos.iloc[1]["_SIMILITUD"] if len(candidatos) > 1 else 0
    )
    confiable = mejor["_SIMILITUD"] >= 0.90 and (
        mejor["_SIMILITUD"] >= 0.95
        or mejor["_SIMILITUD"] - segundo_puntaje >= 0.10
    )
    if confiable:
        return mejor, "NOMBRE APROXIMADO"
    return None, "SIN COINCIDENCIA"


def evaluar_retiro(
    relacion,
    fecha_nacimiento,
    fecha_fallecimiento,
    fecha_retiro,
):
    if pd.isna(fecha_retiro):
        return np.nan, "", "ALERTA: Falta fecha de finalizacion"

    if relacion == "H":
        if pd.isna(fecha_nacimiento):
            return "H", "", "ALERTA: Falta fecha de nacimiento"
        # Se conserva el divisor usado por el archivo de revision oficial.
        edad_retiro = (fecha_retiro - fecha_nacimiento).days / 365.2
        edad_entera_retiro = edad_cumplida(
            fecha_nacimiento, fecha_retiro
        )
        fecha_18 = sumar_anos(fecha_nacimiento, 18)

        # El archivo revisado acepta el retiro durante los 21 años. También
        # contiene casos históricos aprobados cuya fecha final fue exactamente
        # el cumpleaños 18. Esta excepción no acepta retiros a los 19 o 20.
        retiro_historico_18 = fecha_retiro.normalize() == fecha_18.normalize()
        if edad_entera_retiro == 21 or retiro_historico_18:
            resultado = "OK"
        elif edad_entera_retiro < 21:
            resultado = (
                "ALERTA: Hijo retirado antes de la edad permitida "
                f"({edad_retiro:.2f} años)"
            )
        else:
            resultado = (
                "ALERTA: Hijo retirado después de los 21 años "
                f"({edad_retiro:.2f} años)"
            )
        return "H", edad_retiro, resultado

    if relacion == "C":
        if pd.isna(fecha_nacimiento) or pd.isna(fecha_fallecimiento):
            return np.nan, "", "ALERTA: Faltan fechas para la conyuge"

        edad_exacta = (
            fecha_fallecimiento - fecha_nacimiento
        ).days / 365.2
        edad_entera = edad_cumplida(
            fecha_nacimiento, fecha_fallecimiento
        )

        if edad_entera <= 50:
            fecha_correcta = sumar_anos(fecha_fallecimiento, 5)
            validacion = fecha_correcta
        elif edad_entera <= 55:
            fecha_correcta = sumar_anos(fecha_fallecimiento, 6)
            validacion = fecha_correcta
        else:
            return (
                edad_exacta,
                "VITALICIA",
                "ALERTA: Cónyuge con renta vitalicia fue retirada",
            )

        diferencia_dias = (fecha_retiro - fecha_correcta).days
        if diferencia_dias < 0:
            resultado = (
                "ALERTA: Cónyuge retirada antes del plazo "
                f"({abs(diferencia_dias)} días)"
            )
        else:
            resultado = "OK"
        return edad_exacta, validacion, resultado

    return np.nan, "", f"ALERTA: Relacion no contemplada ({relacion})"


def ajustar_formato(archivo, cantidad_filas):
    libro = load_workbook(archivo)
    hoja = libro["Retiradas"]

    hoja.merge_cells("A1:U1")
    hoja.merge_cells("A2:U2")
    hoja.merge_cells("Q4:R4")
    hoja["A1"] = TITULO_PERIODO
    hoja["A2"] = (
        "Validación Pago Final Beneficiarios Pensiones Sobrevivencia"
    )
    hoja["Q4"] = "Revisión"

    azul = PatternFill("solid", fgColor="1F4E78")
    azul_claro = PatternFill("solid", fgColor="D9EAF7")
    rojo = PatternFill("solid", fgColor="FFC7CE")
    verde = PatternFill("solid", fgColor="C6EFCE")

    for celda in hoja[1]:
        celda.fill = azul
        celda.font = Font(color="FFFFFF", bold=True, size=14)
    for celda in hoja[2]:
        celda.fill = azul
        celda.font = Font(color="FFFFFF", italic=True)
    for celda in hoja[4]:
        celda.fill = azul_claro
        celda.font = Font(bold=True)
        celda.alignment = Alignment(horizontal="center")
    for celda in hoja[5]:
        celda.fill = azul_claro
        celda.font = Font(bold=True)
        celda.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )

    for fila in range(6, cantidad_filas + 6):
        for columna in (7, 9, 12, 14, 20):
            hoja.cell(fila, columna).number_format = "dd/mm/yyyy"
        hoja.cell(fila, 13).number_format = "#,##0.00"
        hoja.cell(fila, 16).number_format = "0.00"
        if hoja.cell(fila, 10).value == "H":
            hoja.cell(fila, 17).number_format = "0.000000"
        elif hoja.cell(fila, 17).value != "VITALICIA":
            hoja.cell(fila, 17).number_format = "dd/mm/yyyy"
        resultado = str(hoja.cell(fila, 18).value)
        hoja.cell(fila, 18).fill = (
            rojo if resultado.startswith("ALERTA") else verde
        )

    anchos = {
        1: 7, 2: 20, 3: 20, 4: 8, 5: 23, 6: 38, 7: 18,
        8: 38, 9: 19, 10: 18, 11: 16, 12: 16, 13: 16, 14: 20,
        15: 14, 16: 20, 17: 25, 18: 48, 19: 24, 20: 20, 21: 42,
    }
    for numero, ancho in anchos.items():
        hoja.column_dimensions[get_column_letter(numero)].width = ancho

    hoja.freeze_panes = "A6"
    hoja.auto_filter.ref = f"A5:U{cantidad_filas + 5}"
    hoja.row_dimensions[5].height = 55
    hoja.sheet_view.showGridLines = False
    libro.save(archivo)


def main():
    nomina = pd.read_excel(
        ARCHIVO_NOMINA,
        sheet_name=HOJA_NOMINA,
        engine="openpyxl",
        dtype={"CEDULA_AFILIADO": str, "NUMERO_SOLICITUD": str},
    )
    hojas_base = pd.read_excel(
        ARCHIVO_BASE,
        sheet_name=None,
        engine="openpyxl",
        dtype={"CEDULA": str, "NUMERO_SOLICITUD": str},
    )
    base = pd.concat(
        [
            hoja.assign(Hoja_origen=nombre)
            for nombre, hoja in hojas_base.items()
        ],
        ignore_index=True,
    )

    # Se usan posiciones estables del formato oficial de Nomina para tolerar
    # encabezados con tildes danadas por la codificacion del archivo.
    columnas_nomina = {
        "COMPANIA": nomina.columns[0],
        "SOLICITUD": nomina.columns[1],
        "AFP": nomina.columns[2],
        "CEDULA": nomina.columns[3],
        "NOMBRES_AFILIADO": nomina.columns[4],
        "FECHA_FALLECIMIENTO": nomina.columns[5],
        "DOCUMENTO_BENEFICIARIO": nomina.columns[6],
        "NOMBRES_BENEFICIARIO": nomina.columns[7],
        "FECHA_NACIMIENTO_BENEFICIARIO": nomina.columns[9],
        "RELACION": nomina.columns[10],
        "DISCAPACITADO": nomina.columns[12],
        "ESTATUS": nomina.columns[13],
        "FECHA_PAGO": nomina.columns[15],
        "RENTA_NETA": nomina.columns[19],
        "FECHA_FIN": nomina.columns[21],
        "TIPO_OPERACION": nomina.columns[27],
    }
    nomina = nomina.rename(
        columns={valor: clave for clave, valor in columnas_nomina.items()}
    )

    retiradas = nomina.loc[
        nomina["ESTATUS"].astype(str).str.strip().str.upper().eq("R")
    ].copy()

    base["NOMBRE_BENEFICIARIO"] = nombre_beneficiario_base(base)
    base["_CEDULA"] = base["CEDULA"].map(normalizar_cedula)
    base["_SOLICITUD"] = base["NUMERO_SOLICITUD"].map(
        normalizar_solicitud
    )
    base["_DOCUMENTO_BEN"] = base["ID_BENEFICIARIO"].map(
        normalizar_documento
    )
    base["_NOMBRE_BEN"] = base["NOMBRE_BENEFICIARIO"].map(
        texto_comparable
    )
    base["FECHA_TRANSMISION"] = pd.to_datetime(
        base["FECHA_TRANSMISION"], errors="coerce"
    )
    base.sort_values("FECHA_TRANSMISION", ascending=False, inplace=True)

    retiradas["_CEDULA"] = retiradas["CEDULA"].map(normalizar_cedula)
    retiradas["_SOLICITUD"] = retiradas["SOLICITUD"].map(
        normalizar_solicitud
    )
    retiradas["_DOCUMENTO_BEN"] = retiradas[
        "DOCUMENTO_BENEFICIARIO"
    ].map(normalizar_documento)
    retiradas["_NOMBRE_BEN"] = retiradas["NOMBRES_BENEFICIARIO"].map(
        texto_comparable
    )

    filas_salida = []
    conteo_cruces = {}

    for numero, (_, fila) in enumerate(retiradas.iterrows(), start=1):
        registro_base, metodo_cruce = seleccionar_registro_base(fila, base)
        conteo_cruces[metodo_cruce] = conteo_cruces.get(metodo_cruce, 0) + 1

        fecha_pago = pd.to_datetime(
            fila["FECHA_PAGO"], errors="coerce", dayfirst=True
        )
        if registro_base is None:
            fecha_fin = pd.to_datetime(
                fila["FECHA_FIN"], errors="coerce", dayfirst=True
            )
            fecha_fallecimiento = pd.NaT
            fecha_nacimiento = pd.NaT
            relacion = ""
            discapacitado = ""
            nombre_beneficiario = fila["NOMBRES_BENEFICIARIO"]
            edad_fallecimiento = np.nan
            validacion = ""
            resultado = "ALERTA: Sin coincidencia en la base de datos"
        else:
            # La fecha de finalizacion validada viene de la base. En Nomina
            # hay fechas con dia y mes invertidos que producen alertas falsas.
            fecha_fin = pd.to_datetime(
                registro_base["FECHA_FIN_RENTA"], errors="coerce"
            )
            fecha_fallecimiento = pd.to_datetime(
                registro_base["FECHA_FALLECIMIENTO"], errors="coerce"
            )
            fecha_nacimiento = pd.to_datetime(
                registro_base["BEN_FECHA_NACIMIENTO"], errors="coerce"
            )
            relacion = str(registro_base["BEN_RELACION"]).strip().upper()
            discapacitado = (
                "" if pd.isna(registro_base["DISCAPACITADO"])
                else str(registro_base["DISCAPACITADO"]).strip().upper()
            )
            nombre_beneficiario = registro_base["NOMBRE_BENEFICIARIO"]
            edad_fallecimiento, validacion, resultado = evaluar_retiro(
                relacion,
                fecha_nacimiento,
                fecha_fallecimiento,
                fecha_fin,
            )

        filas_salida.append(
            {
                "No.": numero,
                "Compañía de Seguros": fila["COMPANIA"],
                "Número de Solicitud": fila["SOLICITUD"],
                "AFP": fila["AFP"],
                "Cédula de Identidad del Afiliado": normalizar_cedula(
                    fila["CEDULA"]
                ),
                "Nombres y Apellidos Afiliado": fila["NOMBRES_AFILIADO"],
                "Fecha Fallec. del Afiliado": fecha_fallecimiento,
                "Nombres y Apellidos Beneficiario": nombre_beneficiario,
                "Fecha Nacimiento Benef.": fecha_nacimiento,
                "Relación Benef. con Afiliado": relacion,
                "Hijo Discapac.": discapacitado,
                "Fecha de Pago": fecha_pago,
                "Renta Neta Mes": pd.to_numeric(
                    fila["RENTA_NETA"], errors="coerce"
                ),
                "Fecha Finalización Renta": fecha_fin,
                "Tipo Operac.": fila["TIPO_OPERACION"],
                "Edad a Fallecimiento": edad_fallecimiento,
                "Validación H - Edad Retiro / C - Fecha Retiro": validacion,
                "Resultado revisión automática": resultado,
                "Responsable de Revisión": "",
                "Fecha de Revisión": pd.NaT,
                "Observación": (
                    "" if metodo_cruce in {"DOCUMENTO", "NOMBRE"}
                    else f"Cruce: {metodo_cruce}"
                ),
            }
        )

    salida = pd.DataFrame(filas_salida)
    with pd.ExcelWriter(ARCHIVO_SALIDA, engine="openpyxl") as escritor:
        salida.to_excel(
            escritor,
            sheet_name="Retiradas",
            startrow=4,
            index=False,
        )
    ajustar_formato(ARCHIVO_SALIDA, len(salida))

    alertas = salida["Resultado revisión automática"].str.startswith(
        "ALERTA", na=False
    ).sum()
    print(f"Retiradas encontradas en Nomina: {len(retiradas)}")
    print(f"Cruces con la base: {conteo_cruces}")
    print(f"Registros OK: {len(salida) - alertas}")
    print(f"Alertas: {alertas}")
    print(f"Archivo generado: {ARCHIVO_SALIDA}")


if __name__ == "__main__":
    main()
