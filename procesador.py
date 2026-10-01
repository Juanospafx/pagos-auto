import re
import unicodedata
from datetime import datetime
from difflib import SequenceMatcher
from io import BytesIO

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def normalizar_texto(valor):
    if pd.isna(valor):
        return ""
    texto = unicodedata.normalize("NFKD", str(valor).upper())
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    return re.sub(r"[^A-Z0-9]", "", texto)


def normalizar_solicitud(valor):
    return re.sub(r"^SCOT", "CREC", normalizar_texto(valor))


def normalizar_documento(valor, completar=True):
    texto = normalizar_texto(valor)
    original = "" if pd.isna(valor) else str(valor).strip()
    if re.fullmatch(r"\d+\.0", original):
        texto = texto[:-1]
    if texto in {"", "0", "NAN", "NONE"}:
        return ""
    return texto.zfill(11) if completar and texto.isdigit() else texto


def nombre_base(df):
    return (
        df[["BEN_PRIMER_NOMBRE", "BEN_SEGUNDO_NOMBRE", "BEN_PRIMER_APELLIDO", "BEN_SEGUNDO_APELLIDO"]]
        .fillna("")
        .astype(str)
        .agg(" ".join, axis=1)
        .str.upper()
        .str.strip()
        .str.replace(r"\s+", " ", regex=True)
    )


def similitud(a, b):
    return SequenceMatcher(None, normalizar_texto(a), normalizar_texto(b)).ratio()


def tokens_nombre(valor):
    conectores = {"DE", "DEL", "LA", "LAS", "LOS", "Y"}
    texto = unicodedata.normalize("NFKD", str(valor).upper())
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    return {
        token for token in re.findall(r"[A-Z0-9]+", texto)
        if token not in conectores
    }


def nombres_compatibles(a, b):
    """Acepta cambios de orden y apellidos adicionales sin unir homonimos debiles."""
    tokens_a = tokens_nombre(a)
    tokens_b = tokens_nombre(b)
    if len(tokens_a) < 2 or len(tokens_b) < 2:
        return False
    return tokens_a.issubset(tokens_b) or tokens_b.issubset(tokens_a)


def leer_archivo(contenido, hoja=None, todas=False):
    archivo = BytesIO(contenido)
    return pd.read_excel(
        archivo,
        sheet_name=None if todas else hoja,
        engine="openpyxl",
        dtype=str,
    )


def preparar_base(contenido):
    hojas = leer_archivo(contenido, todas=True)
    base = pd.concat(
        [df.assign(Hoja_origen=nombre) for nombre, df in hojas.items()],
        ignore_index=True,
    )
    base["NOMBRE_BENEFICIARIO"] = nombre_base(base)
    base["_CEDULA"] = base["CEDULA"].map(normalizar_documento)
    base["_SOLICITUD"] = base["NUMERO_SOLICITUD"].map(normalizar_solicitud)
    base["_DOCUMENTO_BEN"] = base["ID_BENEFICIARIO"].map(normalizar_documento)
    base["_NOMBRE_BEN"] = base["NOMBRE_BENEFICIARIO"].map(normalizar_texto)
    base["FECHA_TRANSMISION"] = pd.to_datetime(base["FECHA_TRANSMISION"], errors="coerce")
    base.sort_values("FECHA_TRANSMISION", ascending=False, inplace=True)
    return base


def preparar_nomina(contenido):
    nomina = leer_archivo(contenido, hoja="NOMINA")
    if len(nomina.columns) < 28:
        raise ValueError("La hoja NOMINA no tiene el formato esperado.")
    nombres = {
        "COMPANIA": nomina.columns[0], "SOLICITUD": nomina.columns[1],
        "AFP": nomina.columns[2], "CEDULA": nomina.columns[3],
        "NOMBRE_AFILIADO": nomina.columns[4], "FECHA_FALLECIMIENTO": nomina.columns[5],
        "DOCUMENTO_BEN": nomina.columns[6], "NOMBRE_BEN": nomina.columns[7],
        "FECHA_NACIMIENTO_BEN": nomina.columns[9], "RELACION": nomina.columns[10],
        "DISCAPACITADO": nomina.columns[12], "ESTATUS": nomina.columns[13],
        "FECHA_PAGO": nomina.columns[15], "RENTA_NETA": nomina.columns[19],
        "FECHA_FIN": nomina.columns[21], "MONTO_PRIMER_PAGO": nomina.columns[25],
        "TIPO_OPERACION": nomina.columns[27],
    }
    nomina = nomina.rename(columns={valor: clave for clave, valor in nombres.items()})
    nomina["_CEDULA"] = nomina["CEDULA"].map(normalizar_documento)
    nomina["_SOLICITUD"] = nomina["SOLICITUD"].map(normalizar_solicitud)
    nomina["_DOCUMENTO_BEN"] = nomina["DOCUMENTO_BEN"].map(normalizar_documento)
    nomina["_NOMBRE_BEN"] = nomina["NOMBRE_BEN"].map(normalizar_texto)
    return nomina


def seleccionar_candidato(fila, tabla, nombre_fila, nombre_tabla):
    cedula = fila.get("_CEDULA", "")
    solicitud = fila.get("_SOLICITUD", "")
    alcances = [
        tabla.loc[tabla["_CEDULA"].eq(cedula) & tabla["_SOLICITUD"].eq(solicitud)],
        tabla.loc[tabla["_SOLICITUD"].eq(solicitud)],
        tabla.loc[tabla["_CEDULA"].eq(cedula)],
    ]
    documento = fila.get("_DOCUMENTO_BEN", "")
    nombre_original = fila.get(nombre_fila, "")
    nombre = normalizar_texto(nombre_original)

    for alcance in alcances:
        candidatos = alcance.copy()
        if candidatos.empty:
            continue

        if documento and "_DOCUMENTO_BEN" in candidatos:
            por_documento = candidatos[candidatos["_DOCUMENTO_BEN"].eq(documento)]
            if not por_documento.empty:
                return por_documento.iloc[0]

        por_nombre = candidatos[candidatos[nombre_tabla].map(normalizar_texto).eq(nombre)]
        if not por_nombre.empty:
            return por_nombre.iloc[0]

        por_tokens = candidatos[
            candidatos[nombre_tabla].map(lambda x: nombres_compatibles(nombre_original, x))
        ]
        if not por_tokens.empty:
            return por_tokens.iloc[0]

        candidatos["_SIMILITUD"] = candidatos[nombre_tabla].map(
            lambda x: similitud(nombre_original, x)
        )
        candidatos.sort_values("_SIMILITUD", ascending=False, inplace=True)
        mejor = candidatos.iloc[0]
        segundo = candidatos.iloc[1]["_SIMILITUD"] if len(candidatos) > 1 else 0
        if mejor["_SIMILITUD"] >= 0.90 and (
            mejor["_SIMILITUD"] >= 0.95 or mejor["_SIMILITUD"] - segundo >= 0.10
        ):
            return mejor
    return None


def fecha(valor):
    if pd.isna(valor):
        return pd.NaT
    if isinstance(valor, (pd.Timestamp, datetime)):
        return pd.Timestamp(valor)
    texto = str(valor).strip()
    if re.match(r"^\d{4}-\d{1,2}-\d{1,2}", texto):
        return pd.to_datetime(texto, errors="coerce", yearfirst=True)
    return pd.to_datetime(texto, errors="coerce", dayfirst=True)


def edad_cumplida(nacimiento, evaluacion):
    if pd.isna(nacimiento) or pd.isna(evaluacion):
        return np.nan
    edad = evaluacion.year - nacimiento.year
    if (evaluacion.month, evaluacion.day) < (nacimiento.month, nacimiento.day):
        edad -= 1
    return edad


def cantidad_indexaciones(inicio, corte):
    if pd.isna(inicio) or pd.isna(corte) or corte < inicio:
        return 0
    anos = corte.year - inicio.year
    if (corte.month, corte.day) < (inicio.month, inicio.day):
        anos -= 1
    return max(anos // 2, 0)


def edad_o_fin(registro):
    nacimiento = fecha(registro.get("BEN_FECHA_NACIMIENTO"))
    fin = fecha(registro.get("FECHA_FIN_RENTA"))
    relacion = str(registro.get("BEN_RELACION", "")).strip().upper()
    if relacion == "H":
        return edad_cumplida(nacimiento, fin)
    if relacion == "C":
        return "VITALICIA" if pd.isna(fin) else fin
    return ""


def aplicar_estilo_tabla(
    buffer,
    hoja,
    columnas_fecha=(),
    columna_validacion=10,
    columna_status=14,
):
    buffer.seek(0)
    libro = load_workbook(buffer)
    ws = libro[hoja]
    azul = PatternFill("solid", fgColor="D9E7F3")
    verde = PatternFill("solid", fgColor="00B050")
    rojo = PatternFill("solid", fgColor="FF5B5B")
    amarillo = PatternFill("solid", fgColor="FFF2CC")
    for celda in ws[1]:
        celda.fill = azul
        celda.font = Font(bold=True)
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 42
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col in range(1, ws.max_column + 1):
        maximo = max(len(str(ws.cell(fila, col).value or "")) for fila in range(1, min(ws.max_row, 80) + 1))
        ws.column_dimensions[get_column_letter(col)].width = min(max(maximo + 2, 12), 34)
    for fila in range(2, ws.max_row + 1):
        for col in columnas_fecha:
            ws.cell(fila, col).number_format = "dd/mm/yyyy"
        for col in range(1, ws.max_column + 1):
            encabezado = str(ws.cell(1, col).value or "")
            if encabezado.startswith("MONTO") or encabezado in {"TOTAL PENSION", "Diferencia"}:
                ws.cell(fila, col).number_format = "#,##0.00"
        validacion = str(ws.cell(fila, columna_validacion).value or "")
        status = (
            str(ws.cell(fila, columna_status).value or "")
            if columna_status else ""
        )
        if validacion == "OK":
            ws.cell(fila, columna_validacion).fill = verde
        elif validacion in {"E", "SIN DATOS"} or validacion.startswith("ALERTA"):
            ws.cell(fila, columna_validacion).fill = rojo
        if status in {"Pendiente", "Suspendido"}:
            ws.cell(fila, columna_status).fill = amarillo
    salida = BytesIO()
    libro.save(salida)
    salida.seek(0)
    return salida


def generar_pagos(nomina_bytes, altas_bytes, base_bytes, fecha_corte):
    nomina = preparar_nomina(nomina_bytes)
    base = preparar_base(base_bytes)
    altas = leer_archivo(altas_bytes, hoja="Altas")
    if len(altas.columns) < 11:
        raise ValueError("La hoja Altas no tiene el formato esperado.")

    altas = altas.rename(columns={
        altas.columns[0]: "SOLICITUD", altas.columns[3]: "CEDULA",
        altas.columns[5]: "FECHA_FALLECIMIENTO_ALTA",
        altas.columns[7]: "NOMBRE_BEN", altas.columns[8]: "RELACION_ALTA",
        altas.columns[9]: "FECHA_RESPUESTA_ALTA",
        altas.columns[10]: "PRIMER_PAGO_AFP",
        altas.columns[12]: "PENSION_MENSUAL_ALTA",
    })
    altas["_CEDULA"] = altas["CEDULA"].map(normalizar_documento)
    altas["_SOLICITUD"] = altas["SOLICITUD"].map(normalizar_solicitud)
    altas["_NOMBRE_BEN"] = altas["NOMBRE_BEN"].map(normalizar_texto)
    # Se conserva exactamente la clasificacion que funciona en test_2.py:
    # una fila es pendiente si su cedula, tal como viene en Altas, no aparece
    # en la columna de cedula de Nomina.
    cedulas_nomina = set(
        nomina["CEDULA"].dropna().astype(str).str.strip()
    )
    corte = pd.Timestamp(fecha_corte)
    filas = []

    for _, alta in altas.iterrows():
        pendiente = str(alta.get("CEDULA", "")).strip() not in cedulas_nomina
        # En las pendientes PBS se consulta exclusivamente para saber si el
        # beneficiario esta suspendido. Los demas datos siguen viniendo de Altas.
        registro_pbs_estado = (
            seleccionar_candidato(alta, base, "NOMBRE_BEN", "NOMBRE_BENEFICIARIO")
            if pendiente else None
        )
        suspendido = bool(
            registro_pbs_estado is not None
            and str(registro_pbs_estado.get("BEN_ESTATUS", "")).strip().upper() == "S"
        )
        registro_base = (
            None if pendiente
            else seleccionar_candidato(alta, base, "NOMBRE_BEN", "NOMBRE_BENEFICIARIO")
        )
        registro_nomina = (
            None if pendiente
            else seleccionar_candidato(alta, nomina, "NOMBRE_BEN", "NOMBRE_BEN")
        )

        monto_afp = pd.to_numeric(alta.get("PRIMER_PAGO_AFP"), errors="coerce")
        monto_aseguradora = (
            np.nan if registro_nomina is None
            else pd.to_numeric(registro_nomina.get("MONTO_PRIMER_PAGO"), errors="coerce")
        )
        diferencia = (
            np.nan if pd.isna(monto_afp) or pd.isna(monto_aseguradora)
            else round(monto_aseguradora - monto_afp, 2)
        )
        if suspendido:
            validacion, comentario = "S", "Suspendido en PBS"
        elif pendiente:
            validacion, comentario = "P", "Pendiente de pago"
        elif pd.isna(diferencia):
            validacion, comentario = "SIN DATOS", "Falta información para comparar los montos"
        elif abs(diferencia) <= 0.01:
            validacion, comentario = "OK", ""
        else:
            validacion, comentario = "E", "Error en pago"

        registro_datos = registro_base if registro_base is not None else registro_nomina
        if pendiente:
            id_beneficiario = ""
            primer_nombre = str(alta.get("NOMBRE_BEN", "")).strip().split(" ")[0]
            relacion = str(alta.get("RELACION_ALTA", "")).strip().upper()
            total_pension = pd.to_numeric(alta.get("PENSION_MENSUAL_ALTA"), errors="coerce")
            inicio = fecha(alta.get("FECHA_RESPUESTA_ALTA"))
            indexaciones = cantidad_indexaciones(inicio, corte)
            edad_fallecimiento = np.nan
            fin_renta = ""
        elif registro_datos is None:
            id_beneficiario = ""
            primer_nombre = str(alta.get("NOMBRE_BEN", "")).strip().split(" ")[0]
            relacion = str(alta.get("RELACION_ALTA", "")).strip().upper()
            total_pension = np.nan
            inicio = fecha(alta.get("FECHA_RESPUESTA_ALTA"))
            indexaciones = cantidad_indexaciones(inicio, corte)
            edad_fallecimiento = np.nan
            fin_renta = ""
            detalle = "Beneficiario no encontrado en PBS ni en Nomina"
            comentario = (comentario + ". " if comentario else "") + detalle
        elif registro_base is not None:
            id_beneficiario = normalizar_documento(registro_base.get("ID_BENEFICIARIO"))
            primer_nombre = str(registro_base.get("BEN_PRIMER_NOMBRE", "")).strip()
            relacion = str(registro_base.get("BEN_RELACION", "")).strip().upper()
            total_pension = pd.to_numeric(registro_base.get("TOTAL_PENSION"), errors="coerce")
            inicio = fecha(registro_base.get("FECHA_RESP_SEGURO"))
            indexaciones = cantidad_indexaciones(inicio, corte)
            edad_fallecimiento = edad_cumplida(
                fecha(registro_base.get("BEN_FECHA_NACIMIENTO")),
                fecha(registro_base.get("FECHA_FALLECIMIENTO")),
            )
            fin_renta = edad_o_fin(registro_base)
        else:
            id_beneficiario = normalizar_documento(registro_nomina.get("DOCUMENTO_BEN"))
            primer_nombre = str(registro_nomina.get("NOMBRE_BEN", alta.get("NOMBRE_BEN", ""))).strip().split(" ")[0]
            relacion = str(registro_nomina.get("RELACION", alta.get("RELACION_ALTA", ""))).strip().upper()
            total_pension = np.nan
            inicio = fecha(alta.get("FECHA_RESPUESTA_ALTA"))
            indexaciones = cantidad_indexaciones(inicio, corte)
            nacimiento = fecha(registro_nomina.get("FECHA_NACIMIENTO_BEN"))
            fallecimiento = fecha(registro_nomina.get("FECHA_FALLECIMIENTO"))
            edad_fallecimiento = edad_cumplida(nacimiento, fallecimiento)
            fin = fecha(registro_nomina.get("FECHA_FIN"))
            if relacion == "H":
                fin_renta = edad_cumplida(nacimiento, fin)
            elif relacion == "C":
                fin_renta = "VITALICIA" if pd.isna(fin) else fin
            else:
                fin_renta = ""

        filas.append({
            "Número de Solicitud": alta.get("SOLICITUD", ""),
            "Cédula del beneficiario": id_beneficiario,
            "Primer nombre": primer_nombre,
            "Relación con el beneficiario": relacion,
            "TOTAL PENSION": total_pension,
            "Cantidad de indexaciones": indexaciones,
            "MONTO PRIMER PAGO AFP": monto_afp,
            "MONTO PRIMER PAGO ASEGURADORA": monto_aseguradora,
            "Diferencia": diferencia,
            "Validación": validacion,
            "Comentario": comentario,
            "Edad al fallecimiento": edad_fallecimiento,
            "Edad/fecha fin de renta": fin_renta,
            "Status": (
                "Suspendido" if suspendido
                else "Pendiente" if pendiente
                else "No pendiente"
            ),
        })

    resultado = pd.DataFrame(filas)
    # El archivo final es una bandeja de revision: se omiten los pagos OK y
    # se conservan errores, pendientes y filas con informacion incompleta.
    resultado = resultado.loc[resultado["Validación"].ne("OK")].copy()
    resultado.reset_index(drop=True, inplace=True)
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        resultado.to_excel(writer, sheet_name="Pagos", index=False)
    return aplicar_estilo_tabla(buffer, "Pagos"), resultado


def generar_retiradas(nomina_bytes, base_bytes):
    nomina = preparar_nomina(nomina_bytes)
    base = preparar_base(base_bytes)
    retiradas = nomina[nomina["ESTATUS"].astype(str).str.strip().str.upper().eq("R")].copy()
    filas = []

    for numero, (_, fila) in enumerate(retiradas.iterrows(), 1):
        registro = seleccionar_candidato(fila, base, "NOMBRE_BEN", "NOMBRE_BENEFICIARIO")
        if registro is None:
            nacimiento = fallecimiento = fin = pd.NaT
            relacion = discapacitado = ""
            nombre_ben = fila.get("NOMBRE_BEN", "")
            edad_fall, validacion, resultado = np.nan, "", "ALERTA: Sin coincidencia en la base"
        else:
            nacimiento = fecha(registro.get("BEN_FECHA_NACIMIENTO"))
            fallecimiento = fecha(registro.get("FECHA_FALLECIMIENTO"))
            fin = fecha(registro.get("FECHA_FIN_RENTA"))
            relacion = str(registro.get("BEN_RELACION", "")).strip().upper()
            discapacitado = str(registro.get("DISCAPACITADO", "") if not pd.isna(registro.get("DISCAPACITADO")) else "")
            nombre_ben = registro.get("NOMBRE_BENEFICIARIO", "")
            if relacion == "H":
                edad_fin = (fin - nacimiento).days / 365.2 if not pd.isna(fin) and not pd.isna(nacimiento) else np.nan
                edad_entera = edad_cumplida(nacimiento, fin)
                historico_18 = not pd.isna(fin) and fin.normalize() == (nacimiento + pd.DateOffset(years=18)).normalize()
                resultado = "OK" if edad_entera == 21 or historico_18 else "ALERTA: Edad de retiro incorrecta"
                edad_fall, validacion = "H", edad_fin
            elif relacion == "C":
                edad_fall = (fallecimiento - nacimiento).days / 365.2
                edad_entera = edad_cumplida(nacimiento, fallecimiento)
                if edad_entera <= 50:
                    esperada = fallecimiento + pd.DateOffset(years=5)
                elif edad_entera <= 55:
                    esperada = fallecimiento + pd.DateOffset(years=6)
                else:
                    esperada = None
                if esperada is None:
                    validacion, resultado = "VITALICIA", "ALERTA: Renta vitalicia retirada"
                else:
                    validacion = esperada
                    resultado = "OK" if fin >= esperada else "ALERTA: Retiro antes del plazo"
            else:
                edad_fall, validacion, resultado = np.nan, "", "ALERTA: Relación no contemplada"

        filas.append({
            "No.": numero, "Compañía de Seguros": fila.get("COMPANIA", ""),
            "Número de Solicitud": fila.get("SOLICITUD", ""), "AFP": fila.get("AFP", ""),
            "Cédula de Identidad del Afiliado": normalizar_documento(fila.get("CEDULA")),
            "Nombres y Apellidos Afiliado": fila.get("NOMBRE_AFILIADO", ""),
            "Fecha Fallec. del Afiliado": fallecimiento,
            "Nombres y Apellidos Beneficiario": nombre_ben,
            "Fecha Nacimiento Benef.": nacimiento,
            "Relación Benef. con Afiliado": relacion, "Hijo Discapac.": discapacitado,
            "Fecha de Pago": fecha(fila.get("FECHA_PAGO")),
            "Renta Neta Mes": pd.to_numeric(fila.get("RENTA_NETA"), errors="coerce"),
            "Fecha Finalización Renta": fin, "Tipo Operac.": fila.get("TIPO_OPERACION", ""),
            "Edad a Fallecimiento": edad_fall,
            "Validación H - Edad Retiro / C - Fecha Retiro": validacion,
            "Resultado revisión automática": resultado,
            "Responsable de Revisión": "", "Fecha de Revisión": pd.NaT, "Observación": "",
        })

    resultado = pd.DataFrame(filas)
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        resultado.to_excel(writer, sheet_name="Retiradas", index=False)
    salida = aplicar_estilo_tabla(
        buffer,
        "Retiradas",
        columnas_fecha=(7, 9, 12, 14, 20),
        columna_validacion=18,
        columna_status=None,
    )
    return salida, resultado
